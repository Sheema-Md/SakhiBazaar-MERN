import os
import shutil
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def copy_assets():
    # Asset paths generated previously
    src_img1 = r"C:\Users\moham\OneDrive\Desktop\SakhiBazaar_New\women_entrepreneurs.png"
    src_img2 = r"C:\Users\moham\OneDrive\Desktop\SakhiBazaar_New\ai_storefront_mockup.png"
    
    # Fallback paths
    fallback_img1 = r"C:\Users\moham\.gemini\antigravity-ide\brain\76176d46-6652-49db-9714-5e1ede23a031\women_entrepreneurs_1782914358801.png"
    fallback_img2 = r"C:\Users\moham\.gemini\antigravity-ide\brain\76176d46-6652-49db-9714-5e1ede23a031\ai_storefront_mockup_1782914388014.png"
    
    dest_img1 = "women_entrepreneurs.png"
    dest_img2 = "ai_storefront_mockup.png"
    
    try:
        # Check and copy image 1
        if not os.path.exists(dest_img1) or os.path.getsize(dest_img1) == 0:
            if os.path.exists(src_img1):
                shutil.copy2(src_img1, dest_img1)
            elif os.path.exists(fallback_img1):
                shutil.copy2(fallback_img1, dest_img1)
                print(f"Copied {fallback_img1} to {dest_img1}")
                
        # Check and copy image 2
        if not os.path.exists(dest_img2) or os.path.getsize(dest_img2) == 0:
            if os.path.exists(src_img2):
                shutil.copy2(src_img2, dest_img2)
            elif os.path.exists(fallback_img2):
                shutil.copy2(fallback_img2, dest_img2)
                print(f"Copied {fallback_img2} to {dest_img2}")
    except Exception as e:
        print(f"Error copying images: {e}")

def create_presentation():
    # Initialize presentation
    prs = Presentation()
    
    # Set to widescreen (16:9)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Custom Palette
    COLOR_DEEP_BERRY = RGBColor(90, 24, 70)      # #5A1846 - Primary (Strength/Elegance)
    COLOR_WARM_GOLD = RGBColor(212, 175, 55)     # #D4AF37 - Accent (Commerce/Bazaar)
    COLOR_PEACH_ACCENT = RGBColor(255, 110, 80)  # #FF6E50 - Soft highlight
    COLOR_BG_CREAM = RGBColor(253, 251, 247)     # #FDFBF7 - Main background
    COLOR_WHITE = RGBColor(255, 255, 255)        # #FFFFFF - Card fill
    COLOR_CHARCOAL = RGBColor(44, 62, 80)        # #2C3E50 - Main body text
    COLOR_BORDER_LIGHT = RGBColor(235, 230, 225) # Soft border color

    # Helper function to set slide background
    def set_slide_bg(slide, rgb_color):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = rgb_color

    # Helper function to add slide headers (for slides 2-8)
    def add_slide_header(slide, title_text):
        # Left line decoration (thick Berry bar)
        dec_bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.4), Inches(0.12), Inches(0.8)
        )
        dec_bar.fill.solid()
        dec_bar.fill.fore_color.rgb = COLOR_DEEP_BERRY
        dec_bar.line.fill.background()

        # Title text box
        title_box = slide.shapes.add_textbox(Inches(1.0), Inches(0.35), Inches(11.5), Inches(0.9))
        tf = title_box.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0)
        tf.margin_top = Inches(0)
        
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.name = "Georgia"
        p.font.size = Pt(36)
        p.font.bold = True
        p.font.color.rgb = COLOR_DEEP_BERRY

        # Horizontal separator line (Gold)
        sep_line = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.25), Inches(11.733), Inches(0.02)
        )
        sep_line.fill.solid()
        sep_line.fill.fore_color.rgb = COLOR_WARM_GOLD
        sep_line.line.fill.background()

    # Helper function to draw cards
    def add_card(slide, left, top, width, height, title, items):
        # Outer Card
        card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_WHITE
        
        # Subtle light border
        card.line.color.rgb = COLOR_BORDER_LIGHT
        card.line.width = Pt(1.5)

        # Card Text Frame
        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.3)
        tf.margin_right = Inches(0.3)
        tf.margin_bottom = Inches(0.3)

        # Card Title
        p_title = tf.paragraphs[0]
        p_title.text = title
        p_title.font.name = "Georgia"
        p_title.font.size = Pt(20)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_DEEP_BERRY
        p_title.space_after = Pt(14)

        # Card Items
        for item in items:
            p_item = tf.add_paragraph()
            p_item.space_after = Pt(8)
            p_item.font.name = "Calibri"
            p_item.font.size = Pt(13)
            
            # Formatting text parts (Support bold markdown syntax '**')
            if "**" in item:
                parts = item.split("**")
                is_bold = False
                if item.startswith("**"):
                    is_bold = True
                
                run_bullet = p_item.add_run()
                run_bullet.text = "•  "
                run_bullet.font.bold = False
                run_bullet.font.color.rgb = COLOR_CHARCOAL
                
                for idx, part in enumerate(parts):
                    if not part and idx == 0:
                        continue
                    run = p_item.add_run()
                    run.text = part
                    run.font.color.rgb = COLOR_CHARCOAL
                    if is_bold:
                        run.font.bold = True
                        if any(k in part for k in ["AI", "Price", "Security", "CDN", "SSO", "Feasibility", "Viability", "Advantages"]):
                            run.font.color.rgb = COLOR_DEEP_BERRY
                    else:
                        run.font.bold = False
                    is_bold = not is_bold
            else:
                run = p_item.add_run()
                run.text = "•  " + item
                run.font.color.rgb = COLOR_CHARCOAL

    slide_layout = prs.slide_layouts[6] # blank layout

    # ==========================================
    # SLIDE 1: Title Slide (Dark Theme)
    # ==========================================
    slide1 = prs.slides.add_slide(slide_layout)
    set_slide_bg(slide1, COLOR_DEEP_BERRY)

    # Decorative visual element - large colored rectangle on the right side
    bg_shape = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(9.5), Inches(0), Inches(3.833), Inches(7.5))
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = RGBColor(75, 12, 46) # darker berry
    bg_shape.line.fill.background()

    # Large Gold Accent Triangle on the divide
    gold_tri = slide1.shapes.add_shape(MSO_SHAPE.RIGHT_TRIANGLE, Inches(8.5), Inches(0), Inches(1.0), Inches(7.5))
    gold_tri.fill.solid()
    gold_tri.fill.fore_color.rgb = COLOR_WARM_GOLD
    gold_tri.line.fill.background()
    gold_tri.rotation = 180

    # Title & Subtitle box
    title_box = slide1.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(8.0), Inches(4.0))
    tf = title_box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0)
    
    p1 = tf.paragraphs[0]
    p1.text = "Sakhi Bazaar"
    p1.font.name = "Georgia"
    p1.font.size = Pt(64)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_WHITE
    p1.space_after = Pt(8)
    
    p2 = tf.add_paragraph()
    p2.text = "AI-Enabled Marketplace for Women Entrepreneurs"
    p2.font.name = "Calibri"
    p2.font.size = Pt(22)
    p2.font.color.rgb = COLOR_WARM_GOLD
    p2.space_after = Pt(24)

    # Accent Gold Horizontal Bar
    bar = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(4.4), Inches(4.5), Inches(0.06))
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLOR_WARM_GOLD
    bar.line.fill.background()

    p3 = tf.add_paragraph()
    p3.space_before = Pt(20)
    p3.text = "PROJECT REVIEW & STRATEGY SLIDES"
    p3.font.name = "Calibri"
    p3.font.size = Pt(14)
    p3.font.bold = True
    p3.font.color.rgb = COLOR_PEACH_ACCENT

    p4 = tf.add_paragraph()
    p4.text = "Problem > Solution > Features > Architecture > Pros & Viability"
    p4.font.name = "Calibri"
    p4.font.size = Pt(12)
    p4.font.color.rgb = RGBColor(220, 220, 220)

    # ==========================================
    # SLIDE 2: The Problem
    # ==========================================
    slide2 = prs.slides.add_slide(slide_layout)
    set_slide_bg(slide2, COLOR_BG_CREAM)
    add_slide_header(slide2, "1. The Problem")

    problems_digital = [
        "**Technical Literacy Barriers**: Rural and independent women entrepreneurs struggle with complex e-commerce listing panels and website builders.",
        "**Copywriting Friction**: Writing professional, grammatically correct product descriptions and engaging social media text is a significant pain point.",
        "**Store Customization Costs**: Hiring web developers to set up digital storefronts is economically impossible for micro-sellers."
    ]

    problems_market = [
        "**Pricing Disadvantage**: Creators lack access to localized bazaar metrics, often severely underpricing their handcrafted products.",
        "**Identity Verification & Trust**: Independent micro-sellers struggle to establish legitimacy and buyer trust in crowded online spaces.",
        "**Transaction Security**: Managing online credit/debit card collections safely without coding knowledge exposes sellers to hacking risks."
    ]

    add_card(slide2, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8), "Digital Barriers for Micro-Merchants", problems_digital)
    add_card(slide2, Inches(6.9), Inches(1.8), Inches(5.6), Inches(4.8), "Marketplace & Financial Hurdles", problems_market)

    # ==========================================
    # SLIDE 3: The Expected Solution (With Image 1)
    # ==========================================
    slide3 = prs.slides.add_slide(slide_layout)
    set_slide_bg(slide3, COLOR_BG_CREAM)
    add_slide_header(slide3, "2. The Expected Solution")

    solution_items = [
        "**Zero-Code Storefront**: Setup a personalized digital shop in seconds by inputting basic product criteria, with no coding required.",
        "**AI Generative Co-Writer**: Instant storytelling descriptions and emoji-rich captions optimized for WhatsApp, Facebook, or Instagram.",
        "**Dynamic Price Benchmarks**: Calculations of category statistics (Min, Max, Avg) directly inside the listing form.",
        "**Vetted Storefront Badges**: Admin console to verify credentials, giving independent sellers a verified trust badge."
    ]
    add_card(slide3, Inches(0.8), Inches(1.8), Inches(6.0), Inches(4.8), "A Tailored E-Commerce Platform", solution_items)

    # Add illustration image 1
    img1_path = "women_entrepreneurs.png"
    if os.path.exists(img1_path):
        slide3.shapes.add_picture(img1_path, Inches(7.3), Inches(1.8), width=Inches(5.2), height=Inches(4.8))
        img_frame = slide3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7.3), Inches(1.8), Inches(5.2), Inches(4.8))
        img_frame.fill.background()
        img_frame.line.color.rgb = COLOR_WARM_GOLD
        img_frame.line.width = Pt(1.5)
    else:
        placeholder_items = ["Image: 'women_entrepreneurs.png' would load here.", "Please place image file in the project folder."]
        add_card(slide3, Inches(7.3), Inches(1.8), Inches(5.2), Inches(4.8), "Women Empowerment Illustration", placeholder_items)

    # ==========================================
    # SLIDE 4: Core Features (The Product)
    # ==========================================
    slide4 = prs.slides.add_slide(slide_layout)
    set_slide_bg(slide4, COLOR_BG_CREAM)
    add_slide_header(slide4, "3. Platform Features")

    col_w = Inches(2.72)
    col_h = Inches(4.8)
    t_pos = Inches(1.8)

    feat_ai = [
        "**AI Descriptions**: Gemini drafts warm, narrative listings.",
        "**AI Social Captions**: Generates short posts with tags.",
        "**One-Click Clipboard**: Instant copying to clipboard for social sharing."
    ]
    feat_price = [
        "**Category Analytics**: Displays min, max, and avg pricing on forms.",
        "**AI Price advisor**: Prompts Gemini for corridor retail suggestions.",
        "**Transparency Widget**: Embeds price spreads for buyers."
    ]
    feat_sec = [
        "**Firebase SSO**: Quick, secure Google authentication.",
        "**Secure Cookie JWT**: Session tokens saved in HTTPOnly cookies.",
        "**Stripe / Razorpay**: Signature-verified webhook checkout (Planned)."
    ]
    feat_perf = [
        "**Auto CDN Pipeline**: Cloudinary compresses uploads to WebP.",
        "**Low-Bandwidth UX**: Fast catalog rendering on rural networks.",
        "**Decoupled Model**: Media assets bypass MongoDB database."
    ]

    add_card(slide4, Inches(0.8), t_pos, col_w, col_h, "AI Assist", feat_ai)
    add_card(slide4, Inches(3.72), t_pos, col_w, col_h, "Price Intelligence", feat_price)
    add_card(slide4, Inches(6.64), t_pos, col_w, col_h, "Identity & Safety", feat_sec)
    add_card(slide4, Inches(9.56), t_pos, col_w, col_h, "Cloud & CDN Performance", feat_perf)

    # ==========================================
    # SLIDE 5: System Architecture & Data Flow (With Image 2)
    # ==========================================
    slide5 = prs.slides.add_slide(slide_layout)
    set_slide_bg(slide5, COLOR_BG_CREAM)
    add_slide_header(slide5, "4. System Architecture & Flow")

    flow_items = [
        "**Frictionless Auth Flow**: Google SSO token verification offloaded to Firebase SDK. Session JWTs are issued via HTTPOnly cookies to block XSS script access.",
        "**Smart Catalog Upload Flow**: Seller submits product data. Images pipeline directly to Cloudinary edge nodes for compression to WebP. Metadata documents store in MongoDB.",
        "**AI Generative Flow**: Backend compiles listing criteria into structured prompts, securely calling the Google Gemini SDK.",
        "**Verified Payments Flow**: Stripe Checkout session coordinates payment keys, verified on backend via signature-check webhooks."
    ]
    add_card(slide5, Inches(0.8), Inches(1.8), Inches(6.8), Inches(4.8), "Core Data Communication Loops", flow_items)

    # Add illustration image 2
    img2_path = "ai_storefront_mockup.png"
    if os.path.exists(img2_path):
        slide5.shapes.add_picture(img2_path, Inches(8.1), Inches(1.8), width=Inches(4.4), height=Inches(4.8))
        img_frame = slide5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(8.1), Inches(1.8), Inches(4.4), Inches(4.8))
        img_frame.fill.background()
        img_frame.line.color.rgb = COLOR_WARM_GOLD
        img_frame.line.width = Pt(1.5)
    else:
        placeholder_items = ["Image: 'ai_storefront_mockup.png' would load here.", "Please place image file in the project folder."]
        add_card(slide5, Inches(8.1), Inches(1.8), Inches(4.4), Inches(4.8), "AI Storefront Mockup", placeholder_items)

    # ==========================================
    # SLIDE 6: System Flow Diagram (Visual boxes and arrows)
    # ==========================================
    slide6 = prs.slides.add_slide(slide_layout)
    set_slide_bg(slide6, COLOR_BG_CREAM)
    add_slide_header(slide6, "System Flow Diagram")

    # Draw visual flow nodes
    node_w = Inches(3.2)
    node_h = Inches(1.8)
    n_top = Inches(2.2)

    # Node 1: React 19 Client
    n1 = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), n_top, node_w, node_h)
    n1.fill.solid()
    n1.fill.fore_color.rgb = COLOR_WHITE
    n1.line.color.rgb = COLOR_DEEP_BERRY
    n1.line.width = Pt(2)
    tf1 = n1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "Presentation Layer\n(React 19 & Tailwind)"
    p.font.bold = True
    p.font.size = Pt(16)
    p.font.color.rgb = COLOR_DEEP_BERRY
    p.alignment = PP_ALIGN.CENTER
    p2 = tf1.add_paragraph()
    p2.text = "\n• Google Firebase Auth SSO\n• Seller Inventory Console\n• Price Comparison Widget"
    p2.font.size = Pt(11)
    p2.font.color.rgb = COLOR_CHARCOAL
    
    # Node 2: Node.js Express Server
    n2 = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.0), n_top, node_w, node_h)
    n2.fill.solid()
    n2.fill.fore_color.rgb = COLOR_WHITE
    n2.line.color.rgb = COLOR_DEEP_BERRY
    n2.line.width = Pt(2)
    tf2 = n2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "API Gateway Server\n(Node.js & Express)"
    p.font.bold = True
    p.font.size = Pt(16)
    p.font.color.rgb = COLOR_DEEP_BERRY
    p.alignment = PP_ALIGN.CENTER
    p2 = tf2.add_paragraph()
    p2.text = "\n• verifyToken JWT Middleware\n• Cloudinary Image Uploader\n• Gemini Prompt Compiler"
    p2.font.size = Pt(11)
    p2.font.color.rgb = COLOR_CHARCOAL

    # Node 3: Databases & Integrations
    n3 = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.2), Inches(1.8), Inches(3.3), Inches(2.6))
    n3.fill.solid()
    n3.fill.fore_color.rgb = COLOR_WHITE
    n3.line.color.rgb = COLOR_DEEP_BERRY
    n3.line.width = Pt(2)
    tf3 = n3.text_frame
    tf3.word_wrap = True
    p = tf3.paragraphs[0]
    p.text = "Infrastructure & Resources"
    p.font.bold = True
    p.font.size = Pt(16)
    p.font.color.rgb = COLOR_DEEP_BERRY
    p.alignment = PP_ALIGN.CENTER
    p2 = tf3.add_paragraph()
    p2.text = "\n• MongoDB (Mongoose Schema docs)\n• Cloudinary CDN (Image WebP cache)\n• Google Gemini API (AI text models)\n• Stripe / Razorpay (Payment API hooks)"
    p2.font.size = Pt(11)
    p2.font.color.rgb = COLOR_CHARCOAL

    # Add visual arrows
    arrow1 = slide6.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(4.15), Inches(3.0), Inches(0.7), Inches(0.3))
    arrow1.fill.solid()
    arrow1.fill.fore_color.rgb = COLOR_WARM_GOLD
    arrow1.line.fill.background()

    arrow2 = slide6.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(8.35), Inches(3.0), Inches(0.7), Inches(0.3))
    arrow2.fill.solid()
    arrow2.fill.fore_color.rgb = COLOR_WARM_GOLD
    arrow2.line.fill.background()

    # Communication Loop Description Card
    desc_items = [
        "**Frictionless Client Auth Loop**: SSO is processed via Firebase; session JWT token validated on Express server via secure cookies.",
        "**Performance-Optimized Asset Loop**: Products are saved in MongoDB, while heavy images route directly to Cloudinary WebP servers."
    ]
    add_card(slide6, Inches(0.8), Inches(4.7), Inches(11.7), Inches(1.9), "Data & Asset Integration Flow Loops", desc_items)

    # ==========================================
    # SLIDE 7: Advantages & Pros
    # ==========================================
    slide7 = prs.slides.add_slide(slide_layout)
    set_slide_bg(slide7, COLOR_BG_CREAM)
    add_slide_header(slide7, "5. Advantages & Pros")

    merchant_pros = [
        "**Technical Equality**: Zero-code panels allow sellers with limited computer skills to establish competitive online presences.",
        "**Professional Marketing**: Storytelling AI product copy captures buyer focus, leading to improved conversions.",
        "**Price Margin Protection**: Dynamic category metrics protect independent artisans from underpricing their efforts."
    ]

    technical_pros = [
        "**Lightweight DB Footprint**: Shifting media storage to Cloudinary and user verification logs to Firebase offloads database performance.",
        "**Secure Session Handling**: JWTs stored in HTTPOnly cookies block XSS token extraction vectors.",
        "**Low-Bandwidth Mobile Focus**: Auto-conversion of uploads to WebP guarantees fast page loading on rural, unstable network links."
    ]

    add_card(slide7, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8), "Seller & Business Advantages", merchant_pros)
    add_card(slide7, Inches(6.9), Inches(1.8), Inches(5.6), Inches(4.8), "Technical & Architecture Pros", technical_pros)

    # ==========================================
    # SLIDE 8: Feasibility & Viability
    # ==========================================
    slide8 = prs.slides.add_slide(slide_layout)
    set_slide_bg(slide8, COLOR_BG_CREAM)
    add_slide_header(slide8, "6. Feasibility & Viability")

    feasibility = [
        "**Technical Feasibility**: High. The platform utilizes proven Node/Express, React, and MongoDB patterns. Integrations with Firebase SDK, Cloudinary, and Google Gemini API are fully operational (as demonstrated in the MVP Phase 1 release).",
        "**Operational Feasibility**: High. Zero-code shop dashboards eliminate merchant user training. Admin consoles simplify vetting, business credential review, and content flags, keeping platform operations clean."
    ]

    viability = [
        "**Economic Viability**: Highly Sustainable. Leveraging SaaS resources with generous free tiers (Firebase Auth, Cloudinary base storage, Gemini developer key brackets) keeps base hosting costs minimal.",
        "**Market Monetization Model**: The platform scales by taking a minor commission percentage on Stripe/Razorpay checkouts or offering premium verification tier upgrades for sellers."
    ]

    add_card(slide8, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8), "Project Feasibility", feasibility)
    add_card(slide8, Inches(6.9), Inches(1.8), Inches(5.6), Inches(4.8), "Project Economic Viability", viability)

    # Save presentation
    output_filename = "Sakhi_Bazaar_Project_Review_v3.pptx"
    prs.save(output_filename)
    print(f"Presentation saved successfully as '{output_filename}'")

if __name__ == "__main__":
    copy_assets()
    create_presentation()
